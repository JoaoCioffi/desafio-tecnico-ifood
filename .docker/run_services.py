"""
Sobe o Postgres, aplica o schema, carrega a despensa e SEGURA O TERMINAL.
O container vive enquanto o script viver: Ctrl+C derruba tudo.

    python .docker/run_services.py            # sobe e fica no ar com o painel
    python .docker/run_services.py --reset    # DESTROI o volume e recomeca
    python .docker/run_services.py --down     # derruba um container ja rodando

Garantias:
  * o schema usa CREATE TABLE IF NOT EXISTS  -> nunca derruba tabela
  * o ETL usa ON CONFLICT DO NOTHING         -> nunca sobrescreve nem duplica
  * os dados vivem num volume nomeado do Docker, fora do repositorio, entao
    quem clonar recebe um banco vazio e o carrega a partir da planilha

Sobre a saida: o painel e a UNICA coisa impressa. A subida nao imprime log
proprio — o passo atual aparece na linha `passo` do bloco INFRA e some quando
termina. Versoes anteriores imprimiam o log e depois tentavam apaga-lo com
escapes de terminal; a conta de "subir N linhas" errava toda vez que uma
linha embrulhava ou a janela rolava, e sobrava lixo na tela. Sem log, nao ha
o que apagar.
"""

from __future__ import annotations

import argparse
import subprocess
import sys
import threading
import time
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

import openpyxl
import psycopg
from colorama import Fore, Style
from colorama import init as colorama_init

# A regra de conversao de unidade e dominio, nao ETL: o MCP tambem usa.
from domain.unidades import normalizar
from painel import acompanhar

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent
COMPOSE = AQUI / "docker-compose.yaml"
SCHEMA = AQUI / "db.sql"
PLANILHA = RAIZ / "shared" / "despensa_dona_maria.xlsx"

DSN = "host=localhost port=5432 dbname=sabor_da_maria user=admin password=admin"
VOLUME = "sabor-da-maria-pgdata"
CONTAINER = "sabor-da-maria-db"

ABA_DESPENSA = "Despensa"
ABA_PRECOS = "Precos"

D, B, R = Style.DIM, Style.BRIGHT, Style.RESET_ALL
VERDE, VERM, AMAR = Fore.GREEN, Fore.RED, Fore.YELLOW


# --------------------------------------------------------------------------- #
@dataclass
class Estado:
    """O que a subida ja fez. Escrito pela thread do boot, lido pelo painel.

    Atribuicao de atributo e atomica sob o GIL, e o painel so le — nao ha
    secao critica para proteger. Um Lock aqui seria cerimonia sem funcao.
    """

    passo: str = "conectando ao Docker"
    pronto: bool = False
    erro: str = ""
    etl: dict | None = field(default=None)


# --------------------------------------------------------------------------- #
# Docker
# --------------------------------------------------------------------------- #
def compose(*args: str, check: bool = True) -> subprocess.CompletedProcess:
    proc = subprocess.run(
        ["docker", "compose", "-f", str(COMPOSE), *args],
        capture_output=True, text=True, encoding="utf-8", errors="replace",
    )
    if check and proc.returncode != 0:
        raise RuntimeError((proc.stderr or proc.stdout or "docker compose falhou").strip()[:400])
    return proc


def _saude() -> str:
    p = subprocess.run(
        ["docker", "inspect", "--format", "{{.State.Health.Status}}", CONTAINER],
        capture_output=True, text=True,
    )
    return p.stdout.strip() if p.returncode == 0 else "ausente"


def esperar_banco(estado: Estado, timeout: int = 90) -> None:
    limite = time.monotonic() + timeout
    ultimo = ""
    while time.monotonic() < limite:
        try:
            with psycopg.connect(DSN, connect_timeout=3):
                return
        except psycopg.OperationalError as exc:
            ultimo = (str(exc).strip().splitlines() or [""])[0]
            time.sleep(0.5)
    raise RuntimeError(f"banco nao respondeu em {timeout}s — {ultimo}")


# --------------------------------------------------------------------------- #
# Schema
# --------------------------------------------------------------------------- #
def aplicar_schema(conn: psycopg.Connection) -> None:
    if not SCHEMA.is_file():
        raise RuntimeError(f"{SCHEMA} nao encontrado")
    conn.execute(SCHEMA.read_text(encoding="utf-8"))
    conn.commit()


# --------------------------------------------------------------------------- #
# ETL
# --------------------------------------------------------------------------- #
def ler_planilha() -> list[dict]:
    """Junta as duas abas. Nao assume que casam 1:1 — verifica."""
    if not PLANILHA.is_file():
        raise RuntimeError(f"planilha nao encontrada em {PLANILHA}")

    wb = openpyxl.load_workbook(PLANILHA, data_only=True)
    if faltando := {ABA_DESPENSA, ABA_PRECOS} - set(wb.sheetnames):
        raise RuntimeError(f"abas ausentes: {sorted(faltando)}")

    despensa = {n: (q, u) for n, q, u in _linhas(wb[ABA_DESPENSA], 3)}
    precos = {n: (q, u, p) for n, q, u, p in _linhas(wb[ABA_PRECOS], 4)}

    if so_d := sorted(set(despensa) - set(precos)):
        raise RuntimeError(f"so em {ABA_DESPENSA}: {so_d}")
    if so_p := sorted(set(precos) - set(despensa)):
        raise RuntimeError(f"so em {ABA_PRECOS}: {so_p}")
    if div := [n for n in despensa if str(despensa[n][1]).strip() != str(precos[n][1]).strip()]:
        raise RuntimeError(f"unidade diferente entre as abas: {div}")

    itens = []
    for nome, (estoque, unidade) in despensa.items():
        qtd_comprada, _, preco_pago = precos[nome]
        texto = str(unidade).strip()
        base, fator = normalizar(texto)
        itens.append({
            "nome": nome, "unidade": texto, "estoque": _num(estoque),
            "qtd_comprada": _num(qtd_comprada), "preco_pago": _num(preco_pago),
            "unidade_base": base, "fator": fator,
        })
    return itens


def _linhas(aba, n: int):
    for linha in aba.iter_rows(min_row=2, values_only=True):
        if linha and linha[0] and str(linha[0]).strip():
            yield (str(linha[0]).strip(), *linha[1:n])


def _num(v) -> Decimal | None:
    return None if v is None else Decimal(str(v).replace(",", "."))


def carregar(conn: psycopg.Connection, itens: list[dict]) -> dict:
    """Carrega a despensa e devolve as metricas — o painel mostra no INFRA."""
    inicio = time.monotonic()
    antes = _contar(conn)
    with conn.cursor() as cur:
        cur.executemany(
            """
            INSERT INTO ingredientes
                   (nome, unidade, estoque, qtd_comprada, preco_pago, unidade_base, fator)
            VALUES (%(nome)s, %(unidade)s, %(estoque)s, %(qtd_comprada)s,
                    %(preco_pago)s, %(unidade_base)s, %(fator)s)
            ON CONFLICT (nome) DO NOTHING
            """,
            itens,
        )
    conn.commit()
    return {
        "itens": len(itens),
        "novos": _contar(conn) - antes,
        "segundos": time.monotonic() - inicio,
        "convertidos": sum(1 for i in itens if i["fator"] not in (None, 1)),
        "pendentes": [i["nome"] for i in itens if i["fator"] is None],
    }


def _contar(conn: psycopg.Connection) -> int:
    return conn.execute("SELECT count(*) FROM ingredientes").fetchone()[0]


# --------------------------------------------------------------------------- #
# Subida
# --------------------------------------------------------------------------- #
def subir(estado: Estado) -> None:
    """Toda a subida, em thread. So mexe no `estado` — nao imprime nada.

    Qualquer excecao vira `estado.erro` e o painel a mostra: uma thread que
    morre calada deixaria o painel girando para sempre num passo que nao
    avanca mais.
    """
    try:
        estado.passo = "subindo o container do Postgres"
        compose("up", "-d")

        # Container sobrando de uma execucao anterior nao e recriado pelo
        # `up -d`: se ele estiver doente, fica doente, e a espera abaixo
        # falha por 90s sem dizer o porque. Uma recriacao resolve.
        if _saude() == "unhealthy":
            estado.passo = "recriando o container (estava unhealthy)"
            compose("up", "-d", "--force-recreate")

        estado.passo = "aguardando o banco aceitar conexao"
        esperar_banco(estado)

        with psycopg.connect(DSN) as conn:
            estado.passo = "aplicando o schema"
            aplicar_schema(conn)

            estado.passo = "lendo a planilha da despensa"
            itens = ler_planilha()

            estado.passo = f"gravando {len(itens)} ingredientes"
            estado.etl = carregar(conn, itens)

        estado.passo = ""
        estado.pronto = True
    except Exception as e:
        estado.erro = f"{type(e).__name__}: {e}"


# --------------------------------------------------------------------------- #
# Comandos avulsos
# --------------------------------------------------------------------------- #
def resetar() -> None:
    print(f"\n  {AMAR}▲{R} apaga o volume {B}{VOLUME}{R} e {B}todos os dados{R}.")
    print("    a despensa volta da planilha; perfil, pratos e cardapio somem.")
    if input("  › confirma? digite 'sim': ").strip().lower() != "sim":
        raise SystemExit("  cancelado")
    compose("down", "-v", check=False)
    subprocess.run(["docker", "volume", "rm", "-f", VOLUME], capture_output=True)
    print(f"  {VERDE}✓{R} volume removido\n")


def encerrar() -> None:
    print(f"\n  {D}derrubando o container...{R}", flush=True)
    compose("down", check=False)
    print(f"  {VERDE}✓{R} container derrubado, dados preservados no volume\n")


# --------------------------------------------------------------------------- #
def main() -> None:
    ap = argparse.ArgumentParser(description="Sobe o Postgres, aplica o schema e carrega a despensa.")
    ap.add_argument("--reset", action="store_true", help="DESTROI o volume e recomeca do zero")
    ap.add_argument("--down", action="store_true", help="derruba um container ja rodando")
    args = ap.parse_args()

    if args.down:
        encerrar()
        return
    if args.reset:
        resetar()

    estado = Estado()
    threading.Thread(target=subir, args=(estado,), daemon=True).start()
    acompanhar(estado, DSN, CONTAINER)


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    colorama_init(autoreset=False)

    servindo = not any(a in sys.argv for a in ("--down", "-h", "--help"))
    try:
        main()
    except KeyboardInterrupt:
        print()
    finally:
        if servindo:
            encerrar()
